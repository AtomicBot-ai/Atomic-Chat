"""Mirror the upstream gateway into the tracker, and read the user's decisions back.

    python scripts/upstream-gateway-tracker.py export
        Rewrites the tracker's 'Fork features' and 'Upstream impact' sheets from
        docs/upstream-gateway/fork-features.json and upstream-impact.json.
    python scripts/upstream-gateway-tracker.py import
        Reads the yellow 'Your decision' column of 'Upstream impact' back into
        upstream-impact.json, which is what `make upstream-gate` reads.

Uses openpyxl, the tracker's own tool; nothing in the app depends on it.
See docs/decisions/2026-09-13-gate-every-upstream-sync-on-a-fork-features-impact-report.md.
"""

import datetime
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
TRACKER = ROOT / "docs/superpowers/plans/2026-09-08-model-agnostic-atomic-media-platform-tracker.xlsx"
REGISTER = ROOT / "docs/upstream-gateway/fork-features.json"
REPORT = ROOT / "docs/upstream-gateway/upstream-impact.json"

YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEADER = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
DECISIONS = {"keep ours": "keep-ours", "take theirs": "take-theirs", "adapt": "adapt"}
STATUS_WORDS = {"checked": "Checked", "no-check": "NO CHECK YET", "pending": "Waiting for its PR to merge"}


def replace_sheet(workbook, title, headers, widths):
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
        sheet.cell(row=1, column=index).font = HEADER
    sheet.freeze_panes = "A2"
    return sheet


def export():
    register = json.loads(REGISTER.read_text(encoding="utf-8"))
    report = json.loads(REPORT.read_text(encoding="utf-8")) if REPORT.exists() else None
    workbook = openpyxl.load_workbook(TRACKER)

    features = replace_sheet(
        workbook,
        "Fork features",
        ["ID", "Feature", "What it does for you", "Where it lives", "The check that proves it works", "Decision or record", "Status"],
        [26, 30, 50, 50, 50, 34, 22],
    )
    for row in register:
        features.append([
            row["id"],
            row["feature"],
            row["forUsers"],
            ("MUST STAY REMOVED: " if row.get("expectAbsent") else "") + "\n".join(row["paths"]),
            "\n".join(check["run"] for check in row["checks"]) or "none - this is a gap",
            row.get("decision", ""),
            STATUS_WORDS.get(row["status"], row["status"]),
        ])
    for cells in features.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = WRAP

    impact = replace_sheet(
        workbook,
        "Upstream impact",
        ["ID", "Feature", "What upstream does to it", "What this means", "Files", "Upstream commits", "Recommendation", "Your decision", "Decided on"],
        [26, 30, 16, 60, 45, 50, 16, 18, 14],
    )
    if report is None:
        impact.append(["", "No report yet. Run: make upstream-impact"])
    else:
        impact.append([
            "",
            f"Upstream {report['upstream']} at {report['upstreamHead'][:9]}; Radium {report['base']} at {report['baseHead'][:9]}; made {report['generatedAt'][:10]}.",
        ])
        for row in report["rows"]:
            if not row["flagged"]:
                continue
            impact.append([
                row["id"],
                row["feature"],
                row["kind"],
                row["consequence"],
                "\n".join(f"{f['path']} ({f['kind']})" for f in row["files"]),
                "\n".join(f"{c['sha'][:9]} {c['subject']}" for c in row["commits"]),
                (row["recommendation"] or "").replace("-", " "),
                (row["decision"] or "").replace("-", " "),
                row["decidedOn"] or "",
            ])
            impact.cell(row=impact.max_row, column=8).fill = YELLOW
        unowned = report.get("unownedConflicts") or []
        if unowned:
            impact.append(["", "Conflicts in files no feature claims (check none carries Radium work): " + ", ".join(unowned)])
        untouched = [row["feature"] for row in report["rows"] if not row["flagged"]]
        if untouched:
            impact.append(["", "Untouched by this update: " + ", ".join(untouched)])
    for cells in impact.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = WRAP

    workbook.save(TRACKER)
    print(f"Wrote 'Fork features' ({len(register)} rows) and 'Upstream impact' to {TRACKER.name}")


def import_decisions():
    report = json.loads(REPORT.read_text(encoding="utf-8"))
    rows = {row["id"]: row for row in report["rows"]}
    sheet = openpyxl.load_workbook(TRACKER, read_only=True)["Upstream impact"]
    today = datetime.date.today().isoformat()
    changed, problems = 0, []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        feature_id, raw = values[0], values[7] if len(values) > 7 else None
        if not feature_id or feature_id not in rows:
            continue
        text = (raw or "").strip().lower()
        decision = DECISIONS.get(text) if text else None
        if text and decision is None:
            problems.append(f"{feature_id}: '{raw}' is not keep ours, take theirs or adapt")
            continue
        if decision != rows[feature_id]["decision"]:
            rows[feature_id]["decision"] = decision
            rows[feature_id]["decidedOn"] = today if decision else None
            changed += 1
    REPORT.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {changed} decision change(s).")
    for problem in problems:
        print("  PROBLEM:", problem)
    return 1 if problems else 0


if __name__ == "__main__":
    command = sys.argv[1] if len(sys.argv) > 1 else ""
    if command == "export":
        export()
    elif command == "import":
        sys.exit(import_decisions())
    else:
        print(__doc__)
        sys.exit(2)
